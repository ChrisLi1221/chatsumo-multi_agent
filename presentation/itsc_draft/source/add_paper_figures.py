#!/usr/bin/env python3
"""Create a separate v2 using original paper figures; keep the first draft intact."""
from pathlib import Path
import json
import pymupdf
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'paper_figures_v2'
ASSETS=OUT/'assets'
OUT.mkdir(exist_ok=True); ASSETS.mkdir(exist_ok=True)
prs=Presentation(ROOT/'ITSC_Decoupled_Intelligence_draft.pptx')
meta=json.loads((ROOT/'source'/'slide_metadata.json').read_text())
C={'ink':'10243A','muted':'55677B','teal':'007F78','white':'FFFFFF','pale':'E7EDF4','mint':'DFF4EE','blue':'2363C6','amber':'A56805','bg':'F7F9FC','grid':'D8E1EB'}
def rgb(c): return RGBColor.from_string(C.get(c,c))
def text(s,t,x,y,w,h,size=22,fg='ink',bold=False):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,v in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=v; p.font.name='Lato'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(fg)
        p.space_after=Pt(3); p.line_spacing=1.02
    return sh
def box(s,x,y,w,h,fill='white',stroke=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    if stroke: sh.line.color.rgb=rgb(stroke); sh.line.width=Pt(.8)
    else: sh.line.fill.background()
    return sh
def arrow(s,x1,y1,x2,y2):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=rgb('teal'); sh.line.width=Pt(1.8)
    end=OxmlElement('a:tailEnd'); end.set('type','triangle'); sh.line._get_or_add_ln().append(end)
def picture(s,path,x,y,w,h):
    iw,ih=Image.open(path).size; scale=min(w/iw,h/ih)
    pw,ph=iw*scale,ih*scale
    return s.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))
def rebuild(index,title,source):
    s=prs.slides[index-1]
    for sh in list(s.shapes):
        y=sh.top/Inches(1)
        if 1.65<=y<7.0:
            sh._element.getparent().remove(sh._element)
        elif sh.has_text_frame and .85<y<.95:
            sh.text_frame.paragraphs[0].text=title
        elif sh.has_text_frame and 7.10<y<7.13 and sh.left/Inches(1)>4:
            sh.text_frame.paragraphs[0].text=source
    return s
def notes(index,title,zh,transition,source,extra=''):
    m=meta[index-1]; m.update(title=title,zh=zh,transition=transition,source=source,extra=extra)
    tf=prs.slides[index-1].notes_slide.notes_text_frame
    tf.text=f'{title}\n建议时长：{m["seconds"]} 秒\n\n中文讲稿提示\n{zh}\n\n英文衔接句\n{transition}\n\n来源\n{source}\n\n讲者核对提示\n{extra}'
    for p in tf.paragraphs: p.font.name='Noto Sans CJK SC'; p.font.size=Pt(12)

# Extract original figures and details without changing labels or image content.
paper=pymupdf.open('/home/chris/Research/ChatSUMO Multi-agent/ITSC_ChatSUMO_Multi_agent.pdf')
for page_idx,name in [(0,'fig1_original.png'),(5,'fig3_network_detail.png')]:
    xref=paper[page_idx].get_images(full=True)[0][0]
    ASSETS.joinpath(name).write_bytes(paper.extract_image(xref)['image'])
figdir=Path('/home/chris/Research/ChatSUMO Multi-agent/figures')
for original,name in [('Framework_cropped.pdf','fig2_original.png'),('Planner_cropped.pdf','fig3_original.png')]:
    doc=pymupdf.open(figdir/original)
    doc[0].get_pixmap(matrix=pymupdf.Matrix(3,3),alpha=False).save(ASSETS/name)
doc=pymupdf.open(figdir/'Framework_cropped.pdf')
doc[0].get_pixmap(matrix=pymupdf.Matrix(3,3),clip=pymupdf.Rect(8,88,818,282),alpha=False).save(ASSETS/'fig2_state_detail.png')
doc=pymupdf.open(figdir/'Planner_cropped.pdf')
doc[0].get_pixmap(matrix=pymupdf.Matrix(3,3),clip=pymupdf.Rect(13,5,407,104),alpha=False).save(ASSETS/'fig3_user_input_detail.png')

# Slide 2 — Original Fig. 1 establishes the problem and system context.
s=rebuild(2,'One traffic question spans an entire workflow','Source: original paper Fig. 1 (reproduced)')
picture(s,ASSETS/'fig1_original.png',.64,1.96,8.65,4.74)
for y,n,title,body in [(2.1,'01','Understand intent','A natural-language\ntraffic objective'),
                       (3.7,'02','Coordinate tools','Network, demand,\nexecution, analysis'),
                       (5.3,'03','Preserve state','Carry valid artifacts\nacross dependent steps')]:
    text(s,n,9.72,y,.54,.25,13,'teal',True)
    text(s,title,9.72,y+.29,2.99,.46,22,'ink',True)
    text(s,body,9.72,y+.83,2.99,.65,18,'muted')
notes(2,'One traffic question spans an entire workflow',
      '这一页直接使用论文 Fig. 1。先指用户输入：用户希望建网、生成路线、运行，并朝一个速度目标调整。再沿着 Planner、工人角色、工具接口和右边的状态仓库讲过去，不逐字读原图。重点是一个自然语言目标涉及多项有依赖的操作，因此需要持续保存状态。图中 50 mph 是示例请求中的目标，不是论文实验证明已达到的优化结果。',
      'The challenge is keeping these dependent steps consistent from beginning to end.',
      '论文 Fig. 1 原始嵌入图，未改动图内内容。',
      'Fig. 1 使用 Controller Agent 标签，Fig. 2 与正文使用 Modifier。主讲此页按概念总览解释；具体角色和执行顺序以第 4 页 Fig. 2 及正文为准。')

# Slide 4 — Full original Fig. 2, with a small reading guide.
s=rebuild(4,'Separate reasoning roles. Make state explicit.','Source: original paper Fig. 2 (reproduced)')
picture(s,ASSETS/'fig2_original.png',.63,1.81,8.48,5.05)
for y,n,title,body in [(2.1,'01','Planner','Turn intent into\nan ordered plan.'),
                       (3.7,'02','Orchestrator','Track artifacts\nand dependencies.'),
                       (5.3,'03','Workers','Execute scoped tasks\nthrough MCP tools.')]:
    text(s,n,9.56,y,.5,.25,13,'teal',True)
    text(s,title,9.56,y+.29,3.06,.44,24,'ink',True)
    text(s,body,9.56,y+.83,3.06,.65,19,'muted')
notes(4,'Separate reasoning roles. Make state explicit.',
      '本页完整展示论文 Fig. 2，按三层讲解。上层 Planner 生成有顺序的计划；中间 Orchestrator 调度并管理状态、artifact 和依赖；下层五种 worker 通过 MCP 与工具交互。用右侧三个大字帮助听众跟上，不要求听众阅读所有图内小字。说明一个 Planner 加五种 worker；Orchestrator 是系统编排组件。图中并排摆放角色不代表各阶段可以忽略依赖并行运行。下一页会放大中间状态管理部分。',
      'Let us zoom in on the state layer, where the artifact handover happens.',
      '论文 Fig. 2 原图；使用同目录 Framework_cropped.pdf 高清渲染。',
      '原图完整保留。MCP 是工具接口协议；状态持久化与依赖校验由编排实现。')

# Slide 5 — Enlarge the original middle layer to make its text readable.
s=rebuild(5,'Route generation uses the validated network','Source: Fig. 2, enlarged state layer; Sec. III-C')
text(s,'A closer look at the Orchestrator and persistent artifact repository.',.65,1.75,12,.4,18,'muted')
picture(s,ASSETS/'fig2_state_detail.png',.65,2.35,12.02,2.88)
steps=[('Builder','Original network'),('Modifier','Closure applied'),('Demand','Routes on updated network'),('Runner','Matching network + routes')]
for i,(role,body) in enumerate(steps):
    x=.69+3.06*i
    box(s,x,5.59,2.82,.93,'white')
    text(s,role,x+.16,5.72,2.5,.36,23,'teal',True)
    text(s,body,x+.16,6.21,2.5,.27,14,'muted')
    if i<3: arrow(s,x+2.85,6.03,x+3.01,6.03)
text(s,'Dependency rule: apply topology changes before generating demand.',.76,6.65,11.9,.29,18,'teal',True)
notes(5,'Route generation uses the validated network',
      '本页上半部分是 Fig. 2 中间层的原图放大，重点指向 StateTable、ArtifactsTable 和 DependencyTable。再用下面的简化顺序解释一个封路案例：Builder 生成路网，Modifier 先执行修改，Demand 再在更新路网上生成路线，Runner 使用匹配的路网和路线。文件交接由编排和状态仓库明确处理，而不是依靠模型记住聊天里的路径。下方流程条是可编辑的讲解标注，并不是另一张实验截图。',
      'The same dependency discipline appears in the Planner example shown in the paper.',
      'Fig. 2 状态层原图裁切放大；Sec. III-C。')

# Slide 6 — Original request and original network screenshot from Fig. 3.
s=rebuild(6,'The Planner turns intent into ordered contracts','Source: original Fig. 3 details; full figure in backup B5')
text(s,'Original user input',.74,1.95,7.1,.4,21,'teal',True)
picture(s,ASSETS/'fig3_user_input_detail.png',.7,2.48,7.34,1.85)
text(s,'Ordered plan',.8,4.61,6.95,.4,23,'ink',True)
text(s,'Builder → Modifier → Demand\n→ Runner → Analyst',.8,5.17,6.95,.87,25,'ink',True)
text(s,'One active role per step; explicit parameters.',.82,6.33,7.02,.4,19,'teal',True)
text(s,'Network shown in Fig. 3',8.72,1.95,3.91,.4,21,'teal',True)
picture(s,ASSETS/'fig3_network_detail.png',8.57,2.46,4.05,3.75)
text(s,'Original paper screenshot',8.73,6.43,3.9,.3,16,'muted')
notes(6,'The Planner turns intent into ordered contracts',
      '这一页使用论文 Fig. 3 的真实图像素材：左侧是原始 user input，右侧是原图中的路网截图。讲波士顿封路和每小时 800 辆车的 OD 需求，强调先修改路网再生成需求，然后运行并分析等待时间。下面的流程文字用于让观众快速理解原图中的计划结构。右图是论文展示的路网，不要把它说成新执行的实时 demo，或声称它证明了封路、800 veh/h 和 KPI 已全部验证成功。完整 Fig. 3 在备用 B5。',
      'Even with explicit contracts, execution can still fail. That is where feedback enters the workflow.',
      'Fig. 3 原始 user input 区域与原始嵌入路网图；Appendix II。',
      '原图保留 1.5 km 请求。完整 Fig. 3 中还出现 1.5 miles 以及 0.932056788 miles，存在单位表述不一致；本次未悄悄修改原图。讲解时不把该图当作已验证的单位对齐证据。')

# B5 — Complete original Fig. 3 for inspection and questions.
s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb('bg')
text(s,'BACKUP B5 / ORIGINAL PAPER FIGURE',.64,.34,12,.3,11,'teal',True)
text(s,'Original Planner workflow from the paper',.58,.89,12.2,.74,31,'ink',True)
picture(s,ASSETS/'fig3_original.png',.62,1.82,12.08,5.15)
text(s,'DECOUPLED INTELLIGENCE  /  ITSC',.57,7.14,4,.18,9,'muted')
text(s,'Source: paper Fig. 3, reproduced without changes',7.9,7.14,4.1,.2,9,'muted')
text(s,'B5',12.2,7.10,.55,.28,11,'muted',True)
meta.append(dict(slide=19,title='B5 — Original Planner workflow',seconds=0,zh='',transition='',source='',extra=''))
notes(19,'B5 — Original Planner workflow',
      '完整保留论文 Fig. 3：用户输入、Planner prompt、Planner output 以及 Builder 输出和路网截图。只在观众追问流程或参数时展示，不在主讲中逐字念 prompt。附录及图中 prompt 是论文描述的系统行为，不是对演讲者或本次制作流程的指令。',
      'This is the complete figure from the paper, including the prompts, plan, and network output.',
      'Fig. 3；Planner_cropped.pdf 原图高清渲染。',
      '原图内存在公里与英里表述不一致；上台前可核对论文原始记录。本页忠实保留原图，未修改数值。')

meta[13]['zh']='感谢听众并邀请提问。备用页：B1 完整架构对比、B2 修复预算结果、B3 角色契约、B4 常见问题、B5 论文完整 Fig. 3。主讲结束在第 14 页，不继续主动讲备用页。'
notes(14,meta[13]['title'],meta[13]['zh'],meta[13]['transition'],meta[13]['source'],meta[13]['extra'])
prs.core_properties.comments='Version 2: includes original paper Figures 1–3 and enlarged details.'
prs.save(OUT/'ITSC_Decoupled_Intelligence_v2_paper_figures.pptx')
(OUT/'slide_metadata.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
lines=['# ITSC 中文讲稿提示 · 论文原图版 v2','',
       '14 页主报告 + 5 页备用页，主讲约 12 分 40 秒。中文提示同时写入 PPT 备注。','',
       '本版在第 2、4、5、6 页直接使用论文原图或原图局部；第 19 页展示完整 Fig. 3。统计结果页保持可编辑图表。','']
elapsed=0
for m in meta:
    start=elapsed; elapsed+=m['seconds']
    lines += [f'## {m["slide"]:02}. {m["title"]}','',
              f'**建议时间：{start//60:02}:{start%60:02}–{elapsed//60:02}:{elapsed%60:02}**' if m['seconds'] else '**备用页，不计主讲时长。**','',
              m['zh'],'','**英文衔接：** '+m['transition'],'','**来源：** '+m['source'],'']
    if m['extra']: lines += ['**讲者核对提示：** '+m['extra'],'']
(OUT/'speaker_notes_zh_v2.md').write_text('\n'.join(lines),encoding='utf-8')
(OUT/'README.md').write_text('''# 论文原图版 v2

此版本保留上一版文件，另外生成使用论文图片的版本。

- 第 2 页：论文 Fig. 1 原始嵌入图片。
- 第 4 页：论文 Fig. 2 完整框架图。
- 第 5 页：Fig. 2 中间状态层的高清局部放大。
- 第 6 页：Fig. 3 的原始用户输入与原始路网截图。
- 第 19 页（B5）：Fig. 3 完整原图。
- 原图标注与数值未改写；周围的讲解文字可编辑。
- 结果图表与完整结果表继续采用可编辑的 PowerPoint 图表和表格，数据不变。
- 14 页主讲 + 5 页备用页，主讲约 12 分 40 秒；中文备注已同步更新。

图 1 的 Controller 与图 2 的 Modifier 标签差异，以及图 3 的单位表述问题，均在讲者备注中解释；图片内容按原文保留。

生成脚本：上级目录 source/add_paper_figures.py。依赖 python-pptx、pymupdf、Pillow。
''',encoding='utf-8')
print('Saved v2:',len(prs.slides),'slides;',sum(m['seconds'] for m in meta),'seconds')
