#!/usr/bin/env python3
"""Render the exported PDF and inspect PPTX/PDF structure. Requires pymupdf, Pillow, python-pptx."""
from pathlib import Path
import json
import re
import unicodedata
from PIL import Image, ImageDraw
import pymupdf
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
pdf=pymupdf.open(ROOT/'ITSC_Decoupled_Intelligence_draft.pdf')
prs=Presentation(ROOT/'ITSC_Decoupled_Intelligence_draft.pptx')
def norm(t):
    return re.sub(r'\s+','',unicodedata.normalize('NFKC',t)).replace('–','-').replace('—','-')
issues=[]; text_checks=[]; thumbs=[]
for i,(page,slide) in enumerate(zip(pdf,prs.slides)):
    pix=page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False)
    pix.save(ROOT/'preview'/f'slide-{i+1:02}.png')
    im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
    im.thumbnail((640,360))
    cell=Image.new('RGB',(660,395),'#e3e9f0'); cell.paste(im,(10,5))
    ImageDraw.Draw(cell).text((12,371),f'{i+1:02}',fill='#10243a'); thumbs.append(cell)
    rendered=norm(page.get_text())
    for sh in slide.shapes:
        if sh.width<0 or sh.height<0: issues.append([i+1,'negative shape extent',sh.name])
        if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+100 or sh.top+sh.height>prs.slide_height+100:
            issues.append([i+1,'shape outside slide',sh.name])
        if sh.has_text_frame:
            expected=norm(sh.text)
            if expected and expected not in rendered:
                text_checks.append([i+1,sh.text])
    for word in page.get_text('words'):
        if word[0]<-.5 or word[1]<-.5 or word[2]>page.rect.width+.5 or word[3]>page.rect.height+.5:
            issues.append([i+1,'rendered word outside page',word[4]])
for label,subset in [('contact_sheet',thumbs),('overview_01_09',thumbs[:9]),('overview_10_18',thumbs[9:])]:
    out=Image.new('RGB',(1980,((len(subset)+2)//3)*395),'#e3e9f0')
    for j,im in enumerate(subset): out.paste(im,((j%3)*660,(j//3)*395))
    out.save(ROOT/'preview'/f'{label}.jpg',quality=92)
notes=sum(bool(s.notes_slide.notes_text_frame.text) for s in prs.slides)
charts=[]
for i,s in enumerate(prs.slides):
    for sh in s.shapes:
        if sh.has_chart:
            charts.append({'slide':i+1,'series':[{ 'name':se.name,'values':list(se.values)} for se in sh.chart.series]})
report={'pptx_slides':len(prs.slides),'pdf_pages':len(pdf),'nonempty_notes':notes,
        'editable_charts':len(charts),'chart_values':charts,'geometry_issues':issues,
        'text_requiring_visual_review':text_checks}
(ROOT/'source'/'validation_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
if len(pdf)!=18 or len(prs.slides)!=18 or notes!=18 or issues: raise SystemExit(1)
