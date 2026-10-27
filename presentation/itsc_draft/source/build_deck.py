#!/usr/bin/env python3
"""Build the editable ITSC draft from the supplied paper (no external data).

Dependencies: python-pptx. Run from any directory; outputs beside source/.
"""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
W, H = 13.333333, 7.5
C = dict(bg='F7F9FC', white='FFFFFF', ink='10243A', muted='55677B', pale='E7EDF4',
         grid='D8E1EB', teal='007F78', mint='DFF4EE', aqua='64D7C5', blue='2363C6',
         lightblue='EAF1FE', gray='8294AB', amber='A56805', cream='FFF1D6',
         red='A73F49', rose='F9E8EB', navy='10243A', darkpanel='1D354E')
FONT = 'Lato'
prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
prs.core_properties.title = 'Decoupled Intelligence: A Multi-Agent LLM Framework for Controllable Traffic Scenario Generation in SUMO'
prs.core_properties.subject = 'ITSC oral presentation draft; 12–15 minutes'
prs.core_properties.author = 'Shuyang Li; Ruimin Ke'
prs.core_properties.keywords = 'SUMO, LLM, multi-agent, ITSC, traffic simulation'
slides_meta = []

def color(c): return RGBColor.from_string(C.get(c, c))
def rect(s, x, y, w, h, fill='white', stroke=None, radius=False, sw=1):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color(fill)
    if stroke: sh.line.color.rgb = color(stroke); sh.line.width = Pt(sw)
    else: sh.line.fill.background()
    if radius:
        try: sh.adjustments[0] = .12
        except Exception: pass
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    return sh

def text(s, content, x, y, w, h, size=22, fill='ink', bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0, font=FONT):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, line in enumerate(content.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.alignment = align
        p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color(fill)
        p.space_before = Pt(0); p.space_after = Pt(3); p.line_spacing = 1.04
    return sh

def line(s,x1,y1,x2,y2,fill='gray',width=1.5,arrow=False,dash=False):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb=color(fill); sh.line.width=Pt(width)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    if arrow:
        end=OxmlElement('a:tailEnd'); end.set('type','triangle'); end.set('w','sm'); end.set('len','sm')
        sh.line._get_or_add_ln().append(end)
    return sh

def dot(s,x,y,r=.055,fill='teal'):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-r),Inches(y-r),Inches(r*2),Inches(r*2))
    sh.fill.solid(); sh.fill.fore_color.rgb=color(fill); sh.line.fill.background(); return sh

def tag(s,label,x,y,w,fill='mint',fg='teal',size=11):
    rect(s,x,y,w,.34,fill,radius=True)
    text(s,label,x+.1,y+.055,w-.2,.25,size,fg,True)

def base(section,title,subtitle='',source='',dark=False,backup=False):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb=color('navy' if dark else 'bg')
    fg='white' if dark else 'ink'; muted='A9BCD0' if dark else 'muted'
    rect(s,.55,.43,.28,.045,'aqua' if dark else 'teal')
    text(s,section.upper(),.95,.33,10,.3,11,'aqua' if dark else 'teal',True)
    text(s,title,.55,.88,12.25,.95,31,fg,True)
    if subtitle: text(s,subtitle,.58,1.75,12.1,.47,17,muted)
    line(s,.55,7.02,12.78,7.02,'darkpanel' if dark else 'grid',.7)
    text(s,'DECOUPLED INTELLIGENCE  /  ITSC',.57,7.14,4,.18,9,muted)
    if source: text(s,source,4.3,7.115,7.7,.28,9,muted,align=PP_ALIGN.RIGHT)
    n=len(prs.slides)
    text(s,('B'+str(n-14)) if backup else f'{n:02}',12.2,7.095,.55,.28,11,muted,True,align=PP_ALIGN.RIGHT)
    return s

def panel(s,x,y,w,h,title,body,accent='teal',body_size=19):
    rect(s,x,y,w,h,'white',radius=True)
    rect(s,x,y,.055,h,accent)
    text(s,title,x+.22,y+.21,w-.44,.5,23,accent,True)
    text(s,body,x+.22,y+.91,w-.44,h-1.05,body_size)

def note(s,title,seconds,zh,transition,source,extra=''):
    idx=len(prs.slides)
    content=f'{title}\n建议时长：{seconds} 秒\n\n中文讲稿提示\n{zh}\n\n英文衔接句\n{transition}\n\n来源\n{source}'
    if extra: content+='\n\n讲者核对提示（不在投影页显示）\n'+extra
    s.notes_slide.notes_text_frame.text=content
    # Set notes language/font without changing the visible English slides.
    for p in s.notes_slide.notes_text_frame.paragraphs:
        p.font.name='Noto Sans CJK SC'; p.font.size=Pt(12)
    slides_meta.append(dict(slide=idx,title=title,seconds=seconds,zh=zh,transition=transition,source=source,extra=extra))

def legend(s,x,y,items):
    for label,c,w in items:
        rect(s,x,y+.04,.16,.16,c)
        text(s,label,x+.26,y,w,.3,14,'muted')
        x+=w+.65

def chart(s,x,y,w,h,categories,series,maximum,major,percent=False,kind=XL_CHART_TYPE.COLUMN_CLUSTERED):
    data=CategoryChartData(); data.categories=categories
    for name,values,col in series: data.add_series(name,values)
    ch=s.shapes.add_chart(kind,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
    ch.has_legend=False
    ch.has_title=False
    auto=OxmlElement('c:autoTitleDeleted'); auto.set('val','1'); ch._chartSpace.chart.insert(0,auto)
    ch.font.name=FONT; ch.font.size=Pt(14); ch.font.color.rgb=color('muted')
    ch.chart_style=10
    va=ch.value_axis; ca=ch.category_axis
    va.minimum_scale=0; va.maximum_scale=maximum; va.major_unit=major
    va.has_major_gridlines=True
    va.major_gridlines.format.line.color.rgb=color('grid')
    va.major_gridlines.format.line.width=Pt(.65)
    va.tick_labels.font.name=FONT; va.tick_labels.font.size=Pt(12); va.tick_labels.font.color.rgb=color('muted')
    va.tick_labels.number_format='0"%"' if percent else '#,##0'
    va.major_tick_mark=XL_TICK_MARK.NONE; va.minor_tick_mark=XL_TICK_MARK.NONE
    va.format.line.fill.background()
    ca.tick_labels.font.name=FONT; ca.tick_labels.font.size=Pt(15); ca.tick_labels.font.color.rgb=color('ink')
    ca.major_tick_mark=XL_TICK_MARK.NONE; ca.minor_tick_mark=XL_TICK_MARK.NONE
    ca.format.line.color.rgb=color('grid')
    plot=ch.plots[0]
    if hasattr(plot,'gap_width'): plot.gap_width=95
    plot.has_data_labels=True
    dl=plot.data_labels; dl.font.name=FONT; dl.font.size=Pt(15); dl.font.bold=True; dl.font.color.rgb=color('ink')
    dl.position=XL_LABEL_POSITION.OUTSIDE_END if kind==XL_CHART_TYPE.COLUMN_CLUSTERED else XL_LABEL_POSITION.ABOVE
    dl.number_format='0"%"' if percent else '#,##0.0'
    for se,(_,_,col) in zip(ch.series,series):
        se.format.fill.solid(); se.format.fill.fore_color.rgb=color(col)
        se.format.line.color.rgb=color(col)
        if kind!=XL_CHART_TYPE.COLUMN_CLUSTERED:
            from pptx.enum.chart import XL_MARKER_STYLE
            se.format.line.width=Pt(3); se.marker.style=XL_MARKER_STYLE.CIRCLE; se.marker.size=9
            se.marker.format.fill.solid(); se.marker.format.fill.fore_color.rgb=color(col)
            se.marker.format.line.color.rgb=color(col)
    # Transparent chart and plot backgrounds.
    for parent in [ch._chartSpace, ch._chartSpace.chart.plotArea]:
        sp=OxmlElement('c:spPr'); sp.append(OxmlElement('a:noFill'))
        ln=OxmlElement('a:ln'); ln.append(OxmlElement('a:noFill')); sp.append(ln); parent.append(sp)
    return ch

def table(s,x,y,width,headers,rows,colwidths=None,rowheight=.43,font=16):
    shape=s.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(width),Inches(rowheight*(len(rows)+1)))
    tb=shape.table
    if colwidths:
        for col,v in zip(tb.columns,colwidths): col.width=Inches(v)
    for ri,row in enumerate([headers]+rows):
        tb.rows[ri].height=Inches(rowheight)
        for ci,val in enumerate(row):
            cell=tb.cell(ri,ci); cell.text=str(val); cell.margin_left=Inches(.15); cell.margin_right=Inches(.08)
            cell.margin_top=Inches(.06); cell.margin_bottom=Inches(.04)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb=color('navy' if ri==0 else ('white' if ri%2 else 'pale'))
            for p in cell.text_frame.paragraphs:
                p.font.name=FONT; p.font.size=Pt(font); p.font.bold=(ri==0)
                p.font.color.rgb=color('white' if ri==0 else ('teal' if ci==len(row)-1 else 'ink'))
                p.alignment=PP_ALIGN.LEFT if ci<2 else PP_ALIGN.CENTER
    return tb

# 01 — Title
s=base('ITSC  /  Oral presentation','',dark=True)
text(s,'Decoupled\nIntelligence',.62,1.22,8.05,1.72,49,'white',True)
text(s,'A Multi-Agent LLM Framework for\nControllable Traffic Scenario Generation in SUMO',.66,3.21,8.9,1.22,26,'white')
rect(s,.68,4.85,.52,.045,'aqua')
text(s,'Shuyang Li  ·  Ruimin Ke',.68,5.21,8,.45,22,'white',True)
text(s,'Civil and Environmental Engineering\nRensselaer Polytechnic Institute',.68,5.84,8,0.7,17,'A9BCD0')
# Abstract dependency/network motif, drawn with editable vector shapes.
pts=[(10.18,1.2),(11.7,1.72),(9.67,2.72),(11.0,3.07),(12.05,4.12),(9.75,4.77),(11.17,5.63),(12.35,6.23)]
for a,b in [(0,1),(0,2),(1,3),(2,3),(3,4),(2,5),(3,5),(4,6),(5,6),(6,7)]:
    line(s,*pts[a],*pts[b],'darkpanel',5)
for a,b in [(0,2),(2,3),(3,5),(5,6)]: line(s,*pts[a],*pts[b],'aqua',2.5)
for i,(x,y) in enumerate(pts): dot(s,x,y,.09,'aqua' if i in [0,2,3,5,6] else 'gray')
tag(s,'INTENT → EXECUTION',9.66,6.48,2.65,'darkpanel','aqua')
note(s,'Decoupled Intelligence',35,
     '开场先讲研究问题：我们希望用户用自然语言描述一个交通场景，系统就能在 SUMO 中完成建网、生成需求、运行和指标分析。真正的难点不是生成某一份 XML，而是让整个多步骤流程保持一致。这篇工作的核心是把推理职责拆开，并把文件状态交给明确的编排层管理。',
     'Let me start with a traffic question that sounds simple, but requires a surprisingly long workflow.',
     '论文首页：标题、作者、单位；Sec. I。')

# 02 — Motivation
s=base('Motivation','One traffic question becomes five dependent tasks',source='Source: Sec. I; Fig. 3 (adapted example)')
tag(s,'ILLUSTRATIVE USER REQUEST',.65,1.99,3.0)
text(s,'“Close a street in Boston,\nroute 800 vehicles per hour,\nand report average waiting time.”',.68,2.58,6.8,1.68,27,'ink',True)
# Deliberately schematic, with no claim of geographic fidelity.
rect(s,8.1,2.0,4.55,2.52,'white',radius=True)
roads=[(8.42,2.5,12.3,2.5),(8.42,3.12,12.3,3.12),(8.42,3.88,12.3,3.88),
       (8.92,2.25,8.92,4.18),(10.02,2.25,10.02,4.18),(11.35,2.25,11.35,4.18),(12.03,2.25,12.03,4.18)]
for a in roads: line(s,*a,'grid',5)
line(s,10.04,3.12,11.32,3.12,'red',5)
line(s,10.59,2.97,10.81,3.27,'red',2); line(s,10.81,2.97,10.59,3.27,'red',2)
dot(s,8.92,2.5,.1,'teal'); dot(s,12.03,3.88,.1,'blue')
text(s,'O',8.56,2.18,.3,.3,16,'teal',True); text(s,'D',12.2,3.88,.3,.3,16,'blue',True)
text(s,'Road closure • schematic',8.33,4.25,4,.23,11,'muted')
labels=[('01','Build','Road network'),('02','Modify','Street closure'),('03','Demand','O–D flows'),('04','Run','SUMO execution'),('05','Analyze','Waiting time')]
for i,(num,title,body) in enumerate(labels):
    x=.65+i*2.45
    rect(s,x,4.97,2.2,1.23,'white',radius=True)
    text(s,num,x+.16,5.13,.42,.3,13,'teal',True)
    text(s,title,x+.62,5.07,1.45,.45,22,'ink',True)
    text(s,body,x+.17,5.69,1.9,.35,15,'muted')
    if i<4: line(s,x+2.22,5.54,x+2.42,5.54,'gray',1.5,True)
text(s,'Every stage must agree on the network, demand, parameters, and files.',.69,6.54,12,.4,21,'teal',True)
note(s,'One traffic question becomes five dependent tasks',55,
     '用波士顿封路的例子引出问题。对使用者来说，只有一句话；对系统来说，需要先提取路网，再完成封路，然后在修改后的网络上生成每小时 800 辆车的需求，接着运行，最后提取平均等待时间。强调顺序：如果先生成路线再删路，后面就可能出现无效路径。右侧只是示意图，不是波士顿真实地图，也不是实验截图。',
     'The challenge is keeping these dependent steps consistent from beginning to end.',
     'Sec. I；Fig. 3 请求的简化示例。街道图为演示示意。')

# 03 — Failure modes
s=base('Problem','A single context must remember every dependency',
       'The paper’s baseline is a monolithic ReAct agent with all tools and execution history.',
       'Source: Sec. IV-A, IV-C, IV-E')
rect(s,.65,2.49,6.4,3.83,'white',radius=True)
tag(s,'ONE GROWING CONTEXT',.9,2.73,2.9,'lightblue','blue')
chips=[('User goal',.95,3.4,2.05),('Tool calls',3.26,3.4,2.0),('Network path',1.5,4.09,2.37),
       ('Route parameters',4.0,4.09,2.54),('Modified edges',1.0,4.8,2.55),('Retry logs',3.75,4.8,2.1),('KPI outputs',2.4,5.49,2.35)]
for label,x,y,w in chips:
    rect(s,x,y,w,.49,'pale',radius=True); text(s,label,x+.13,y+.09,w-.26,.33,17,'muted')
for i,(a,b) in enumerate([(0,2),(1,3),(2,4),(3,5),(4,6)]):
    ca=chips[a]; cb=chips[b]; line(s,ca[1]+ca[3]/2,ca[2]+.49,cb[1]+cb[3]/2,cb[2],'grid',1)
failure=[('Path reference errors','A downstream step points to the wrong file.'),
         ('Parameter inconsistency','Physical values drift across tool calls.'),
         ('Reasoning drift','Retries lose track of the original objective.')]
for i,(head,body) in enumerate(failure):
    y=2.54+i*1.3
    dot(s,7.65,y+.24,.07,'red')
    text(s,head,7.94,y,4.66,.45,23,'ink',True)
    text(s,body,7.94,y+.54,4.54,.7,18,'muted')
text(s,'Design question: which responsibilities should be explicit in the system?',.68,6.57,12,.35,20,'teal',True)
note(s,'A single context must remember every dependency',50,
     '这里说的是论文实际比较的单体 ReAct 基线。它把所有工具和整个历史放进同一上下文，既要理解任务，又要记文件位置，还要处理封路、路由和重试。论文对失败日志的分析指出路径引用错误与推理漂移。参数不一致也是设计动机之一。不要把这一页扩大成“所有单智能体系统都会失败”的结论；问题是这种具体组织方式面对长流程时的可靠性。',
     'Our response is to separate the reasoning roles and make the execution state explicit.',
     'Sec. IV-A 基线定义；Sec. IV-C / IV-E 失败分析；Sec. I 参数一致性动机。')

# 04 — Architecture
s=base('Method 01 / Role specialization','Separate reasoning roles. Make state explicit.',
       'One Planner + five specialized workers, coordinated through a persistent Orchestrator.',
       'Source: Sec. III; Fig. 2 (redrawn)')
rect(s,3.78,2.39,4.84,.62,'lightblue',radius=True)
text(s,'Planner  ·  goal → ordered plan',3.94,2.53,4.5,.4,21,'blue',True,PP_ALIGN.CENTER)
line(s,6.2,3.03,6.2,3.29,'blue',1.5,True)
rect(s,2.02,3.34,6.77,.86,'navy',radius=True)
text(s,'Orchestrator',2.26,3.49,3.17,.42,25,'white',True)
text(s,'Coordinate • validate • pass files',5.15,3.61,3.39,.29,14,'white')
rect(s,9.34,3.19,3.33,1.12,'mint',radius=True)
text(s,'Persistent state',9.57,3.32,2.91,.35,21,'teal',True)
text(s,'Artifacts + dependencies (SQLite)',9.57,3.85,2.92,.3,13,'teal')
line(s,8.83,3.61,9.28,3.61,'teal',1.4,True)
line(s,9.28,3.94,8.83,3.94,'teal',1.4,True)
line(s,5.4,4.22,5.4,4.48,'gray',1.5,True)
rect(s,.65,4.52,12.03,.44,'pale',radius=True)
text(s,'MODEL CONTEXT PROTOCOL (MCP)  /  TOOL INTERFACE',.86,4.61,11.6,.28,13,'muted',True,PP_ALIGN.CENTER)
workers=[('Builder','OSM → network','teal'),('Modifier','Topology / signals','amber'),('Demand','Flows / routes','teal'),('Runner','SUMO execution','blue'),('Analyst','Metrics / feedback','blue')]
for i,(head,body,col) in enumerate(workers):
    x=.65+i*2.45
    line(s,x+1.1,4.98,x+1.1,5.26,'gray',1.2,True)
    rect(s,x,5.3,2.22,1.02,'white',radius=True)
    text(s,head,x+.17,5.43,1.88,.4,24,col,True,PP_ALIGN.CENTER)
    text(s,body,x+.12,5.96,1.98,.3,14,'muted',align=PP_ALIGN.CENTER)
text(s,'Scoped reasoning',.74,6.6,3.3,.32,18,'teal',True)
text(s,'Validated handovers',4.81,6.6,3.5,.32,18,'teal',True)
text(s,'Bounded feedback',9.12,6.6,3.6,.32,18,'teal',True)
note(s,'Separate reasoning roles. Make state explicit.',80,
     '先从上到下介绍，不要逐字读框。Planner 负责把全局目标拆成有顺序的步骤。Orchestrator 负责执行协调、校验和状态传递；SQLite 保存状态、文件和依赖。MCP 提供统一的工具接口。底部是五个工人角色：Builder 建网、Modifier 修改路网或信号、Demand 生成需求、Runner 执行、Analyst 计算指标并反馈。因此总共是一个 Planner 加五种 worker 角色；Orchestrator 是编排组件。虽然图里并排画了五个角色，论文流程按照依赖顺序执行，不是在并行生成路网和路线。',
     'The key handover is the network artifact: every downstream step must use the validated topology.',
     'Sec. III-A 至 III-C；Fig. 2。',
     '明确区分：MCP 是工具协议；文件状态持久化与依赖校验由 Orchestrator 实现。')

# 05 — State example
s=base('Method 02 / Artifact consistency','Route generation uses the validated network',
       'A road closure changes the physical world that downstream tools must use.',
       'Source: Sec. III-A, III-B, III-C.2–4 (illustrative artifact labels)')
xs=[.72,3.72,6.72,9.72]
items=[('Builder','network v1','Original topology','lightblue','blue'),
       ('Modifier','network v2','Closure applied','cream','amber'),
       ('Demand','routes','Generated on v2','mint','teal'),
       ('Runner','simulation','v2 + matching routes','lightblue','blue')]
for x,(head,file,body,bg,fg) in zip(xs,items):
    rect(s,x,2.63,2.8,2.15,'white',radius=True)
    text(s,head,x+.18,2.84,2.44,.44,24,fg,True)
    rect(s,x+.16,3.55,2.48,.53,bg,radius=True)
    text(s,file,x+.3,3.65,2.2,.32,19,fg,True,PP_ALIGN.CENTER)
    text(s,body,x+.17,4.35,2.5,.3,15,'muted',align=PP_ALIGN.CENTER)
for i in range(3): line(s,xs[i]+2.82,3.66,xs[i+1]-.03,3.66,'gray',1.3,True)
rect(s,.72,5.36,7.13,1.0,'navy',radius=True)
text(s,'Shared state passes the latest validated artifact.',.97,5.6,6.67,.56,22,'white',True)
line(s,5.12,4.82,5.12,5.32,'teal',1.5,True)
text(s,'Dependency rule',8.33,5.38,4.2,.36,19,'teal',True)
text(s,'Modify topology before generating demand.',8.33,5.88,4.16,.65,20)
note(s,'Route generation uses the validated network',60,
     '用一个具体文件交接解释为什么需要状态管理。Builder 得到原始路网，Modifier 执行封路后得到更新路网。Demand 必须拿到更新后的文件，Runner 再把这份路网和对应路线一起放进配置。这样文件引用不是靠模型从聊天记录里回忆，而是由状态仓库明确提供。这里 v1/v2 只是讲解用的逻辑状态标签；论文描述 Modifier 会更新路网 artifact，并没有声称实现了完整文件版本控制。',
     'This dependency rule also appears in how the Planner translates a user request.',
     'Sec. III-A 状态与依赖；Sec. III-C Modifier、Demand、Runner。',
     'v1/v2 是演示标签，不代表实际文件名或版本控制功能。')

# 06 — Concrete planning
s=base('Method / Worked example','The Planner turns intent into ordered contracts',
       'Illustrative Boston request, adapted from the paper’s Fig. 3.',
       'Source: Fig. 3; Appendix II–III (simplified, not an execution log)')
rect(s,.66,2.28,12.0,.76,'navy',radius=True)
text(s,'Boston: close Boylston Street; Newbury Street → Beacon Street at 800 veh/h.\nRun the simulation and report average waiting time.',.9,2.39,11.52,.6,18,'white')
rows=[('01','Builder','Build the Boston road network','.net.xml'),
      ('02','Modifier','Apply the requested Boylston Street closure','Updated network'),
      ('03','Demand','Newbury Street → Beacon Street; 800 veh/h','.rou.xml'),
      ('04','Runner','Configure and execute SUMO','.sumocfg + outputs'),
      ('05','Analyst','Extract average waiting time','KPI report')]
for i,(num,role,task,output) in enumerate(rows):
    y=3.32+i*.59
    rect(s,.66,y,12,.51,'white' if i%2==0 else 'pale',radius=True)
    text(s,num,.88,y+.09,.44,.3,15,'teal',True)
    text(s,role,1.53,y+.075,1.7,.32,18,'ink',True)
    text(s,task,3.44,y+.08,6.0,.34,17)
    text(s,output,10.05,y+.095,2.4,.3,15,'teal')
text(s,'Each step activates one role and supplies only the parameters that role needs.',.73,6.58,11.9,.38,20,'teal',True)
note(s,'The Planner turns intent into ordered contracts',60,
     '沿着这五行快速讲一次：先建 Boston 路网，随后执行 Boylston Street 封路，再根据 Newbury Street 到 Beacon Street 的请求生成每小时 800 辆车的流，最后运行并提取平均等待时间。Planner 的每一步只激活一个角色，并明确必要的输入参数。这个例子展示的是任务分解方式，不是新增的实际运行演示。街道名称需要经过工具解析与验证，不能随便杜撰 edge ID。',
     'Even with explicit contracts, execution can still fail. That is where feedback enters the workflow.',
     'Fig. 3；Appendix II 的单角色步骤约束；Appendix III。',
     '原 Fig. 3 出现 1.5 km、1.5 miles 和 distance_miles=0.932056788 的混合表述。本页省略该尺度字段，避免把原图单位不一致带入主讲。若补现场 demo，先统一尺度含义与单位。')

# 07 — Loops
s=base('Method 03 / Feedback','Two feedback loops address two different problems',
       'Execution repair handles tool failures; KPI feedback revisits the scenario objective.',
       'Source: Sec. III-B; Algorithm 1')
for x,head,body in [(1.0,'Plan','Decompose objective'),(4.5,'Execute','Workers + validation'),(8.0,'Analyze','Extract KPIs')]:
    rect(s,x,2.7,2.73,1.12,'white',radius=True)
    text(s,head,x+.16,2.88,2.41,.4,26,'ink',True,PP_ALIGN.CENTER)
    text(s,body,x+.14,3.43,2.45,.26,14,'muted',align=PP_ALIGN.CENTER)
line(s,3.77,3.24,4.44,3.24,'gray',1.7,True)
line(s,7.27,3.24,7.94,3.24,'gray',1.7,True)
text(s,'valid',7.29,2.9,.66,.22,11,'muted',align=PP_ALIGN.CENTER)
line(s,10.77,3.24,11.35,3.24,'teal',1.7,True)
rect(s,11.42,2.89,1.1,.72,'mint',radius=True); text(s,'Done',11.5,3.06,.94,.35,19,'teal',True,PP_ALIGN.CENTER)
text(s,'target met',10.87,2.51,1.65,.25,12,'teal')
# Distinct return lanes with bounded retries stated below.
line(s,5.87,3.87,5.87,4.65,'amber',1.7)
line(s,5.87,4.65,2.35,4.65,'amber',1.7)
line(s,2.35,4.65,2.35,3.88,'amber',1.7,True)
tag(s,'1  EXECUTION ERROR → STRUCTURED FEEDBACK',.91,4.91,6.2,'cream','amber',13)
line(s,9.37,3.87,9.37,5.72,'blue',1.7)
line(s,9.37,5.72,1.34,5.72,'blue',1.7)
line(s,1.34,5.72,1.34,3.88,'blue',1.7,True)
tag(s,'2  KPI TARGET UNMET → REPLAN',.91,6.02,5.02,'lightblue','blue',13)
text(s,'Bounded attempts\nStop on success or budget limit',9.9,5.29,2.61,1.0,18,'muted')
note(s,'Two feedback loops address two different problems',65,
     '明确区分两条返回路径。第一条是执行级错误，例如文件缺失或路线不可达，系统把结构化错误送回 Planner 重新规划。第二条是在仿真执行有效之后，Analyst 检查指标是否满足用户目标；未达到目标则反馈给 Planner 调整。循环有次数上限。这里展示的是论文提出的双层设计；后面的定量实验主要测的是修复预算对任务成功率的影响，不能把它说成已经量化证明了交通 KPI 的最优性或收敛性。',
     'We evaluate the system first as an execution architecture, and then as a repair mechanism.',
     'Sec. III-B；Algorithm 1；Sec. IV-D。')

# 08 — Setup
s=base('Evaluation / Setup','30 tasks test increasing dependency complexity',
       'SUMO 1.21.0  ·  Qwen3-30B-Coder  ·  monolithic ReAct vs. the proposed framework',
       'Source: Sec. IV-A, IV-B')
for x,lev,head,body,col in [( .67,'L1','Basic','Build a geographic network\n+ qualitative traffic','blue'),
                           (4.77,'L2','Intermediate','Specify origin–destination flow\n+ align tool parameters','teal'),
                           (8.87,'L3','Complex','Modify network topology\nbefore demand generation','amber')]:
    rect(s,x,2.51,3.8,2.37,'white',radius=True)
    tag(s,lev,x+.2,2.73,.57,{'blue':'lightblue','teal':'mint','amber':'cream'}[col],col)
    text(s,head,x+.2,3.31,3.42,.51,25,'ink',True)
    text(s,body,x+.2,4.03,3.42,.72,18,'muted')
text(s,'WHAT IS MEASURED?',.73,5.3,11.8,.3,12,'teal',True)
metrics=[('Success rate ↑','Valid configuration + error-free execution'),('Average tokens ↓','Total LLM input + output tokens'),('Time-to-Insight ↓','User request → first valid KPI')]
for i,(head,body) in enumerate(metrics):
    x=.72+i*4.1
    text(s,head,x,5.89,3.76,.43,22,'ink',True)
    text(s,body,x,6.4,3.78,.48,16,'muted')
note(s,'30 tasks test increasing dependency complexity',55,
     '实验使用 SUMO 1.21.0 和论文所写的 Qwen3-30B-Coder，共 30 个任务。L1 是基本建网和定性需求；L2 需要特定起终点与参数对齐；L3 增加生成需求前的路网结构修改。评价三个量：成功率表示配置有效且运行没有物理错误；平均 token 是全流程输入输出消耗；Time-to-Insight 是从用户输入到拿到第一个有效 KPI 的总墙钟时间。不要把 TTI 说成纯 SUMO 内核运行时间，也不要把成功率说成所有交通目标都已达成。',
     'The largest reliability difference appears in the most complex tasks.',
     'Sec. IV-A / IV-B；任务总数为 30，正文未明确给出各层级样本量。',
     '沿用论文模型名称，未自行改写成另一模型版本。论文未完整报告硬件、推理服务和计时纳入项；绝对时延应限制在所报告设置内解释。')

# 09 — Main success result
s=base('Results / Architecture comparison','Complex-task success rises from 70% to 90%',
       'Success rate by task complexity (%) ↑', 'Source: Table I; Sec. IV-C')
legend(s,.94,2.31,[('Monolithic','gray',1.75),('Multi-agent','teal',1.8)])
chart(s,.67,2.85,7.62,3.37,['L1 · Basic','L2 · Intermediate','L3 · Complex'],
      [('Monolithic',[100,100,70],'gray'),('Multi-agent',[90,100,90],'teal')],110,25,True)
rect(s,8.8,2.39,3.82,3.83,'mint',radius=True)
text(s,'L3 / COMPLEX',9.08,2.66,3.25,.31,13,'teal',True)
text(s,'+20',9.02,3.31,3.3,1.1,65,'teal',True)
text(s,'percentage points',9.08,4.52,3.17,.43,22,'teal',True)
text(s,'The gain is concentrated\nin harder workflows.',9.08,5.23,3.13,.7,18,'ink')
text(s,'L1 is 10 points lower; L2 is unchanged. Role separation helps most when dependencies grow.',.74,6.58,11.97,.35,18,'muted')
note(s,'Complex-task success rises from 70% to 90%',70,
     '先指右侧：在最复杂的 L3，成功率从 70% 提升到 90%，是增加 20 个百分点。然后带观众看左边完整结果：L2 两者都是 100%；L1 单体是 100%，多智能体是 90%。因此这不是每种任务都全面胜出的结果，而是优势主要体现在依赖复杂的任务。可以补充论文失败日志提示路径引用错误与推理漂移，但不要把归因说成独立受控消融已经证明。样本较小且表中没有置信区间，不使用“统计显著”措辞。',
     'The reliability gain is accompanied by lower token use and a shorter time to the first valid result.',
     'Table I；Sec. IV-C。+20 pp = 90% − 70%。',
     '这是整体架构对比，不能拆成“单独增加 Planner / MCP / state 各贡献多少”。L1 降低不能直接用 Table II 的 overthinking 现象作未经验证的因果解释。')

# 10 — Efficiency
s=base('Results / Efficiency','Fewer tokens and a shorter path to the first KPI',
       'Reported averages in the architecture comparison; lower is better.', 'Source: Table I; percentages and ratios calculated from reported values')
legend(s,.89,2.3,[('Monolithic','gray',1.75),('Multi-agent','teal',1.8)])
rect(s,.68,2.86,5.85,3.48,'white',radius=True)
rect(s,6.81,2.86,5.85,3.48,'white',radius=True)
text(s,'Average tokens',.92,3.06,5.27,.34,20,'ink',True)
text(s,'Time-to-Insight (seconds)',7.05,3.06,5.25,.34,20,'ink',True)
ct=chart(s,.76,3.52,5.57,2.64,['L1','L2','L3'],
      [('Monolithic',[8529.2,9690.3,11404.7],'gray'),('Multi-agent',[4845.5,5310.1,6028.7],'teal')],14000,4000)
ct.plots[0].data_labels.number_format='#,##0'; ct.plots[0].data_labels.font.size=Pt(12)
ct=chart(s,6.89,3.52,5.57,2.64,['L1','L2','L3'],
      [('Monolithic',[19.8,22.7,41.4],'gray'),('Multi-agent',[6.2,7.3,8.3],'teal')],50,10)
ct.plots[0].data_labels.font.size=Pt(13)
text(s,'L3: 47.1% fewer tokens',.95,6.53,5.47,.4,23,'teal',True)
text(s,'L3: 41.4 s → 8.3 s  /  ≈5.0× speedup',7.03,6.53,5.48,.4,22,'teal',True)
note(s,'Fewer tokens and a shorter path to the first KPI',70,
     '先讲左图：三个难度下 token 都下降，L3 从 11,404.7 到 6,028.7，减少 47.1%，比直接说“减半”更精确。右图是 Time-to-Insight，L3 从 41.4 秒到 8.3 秒，约五倍加速，也就是时间降低约 80%。论文的解释是角色只接收与子任务有关的上下文，减少冗余推理与无效重试。强调这些是该实验配置下的报告均值，不是对所有模型和部署的速度保证。',
     'Next, we look separately at what happens when the system is allowed to repair its own errors.',
     'Table I。L3 token reduction = (11404.7−6028.7)/11404.7 = 47.14%；TTI ratio = 41.4/8.3 = 4.99。',
     '图中 token 标签为四舍五入整数；完整小数见备用页 B1。TTI 不等同于纯仿真运行耗时。')

# 11 — Repair experiment
s=base('Results / Repair-budget study','One repair attempt helps; a second adds cost',
       'Separate experiment: compare 0, 1, and 2 maximum repair attempts within Table II.',
       'Source: Table II; Sec. IV-D')
text(s,'L2 and L3 success rate (%)',.91,2.39,6.6,.39,22,'ink',True)
chart(s,.78,2.98,6.58,3.07,['0 repairs','1 repair','2 repairs'],
      [('L2 and L3',[30,50,50],'teal')],100,25,True,XL_CHART_TYPE.LINE_MARKERS)
text(s,'L2 and L3 have identical reported success rates.',.94,6.15,6.34,.34,15,'muted')
rect(s,7.89,2.39,4.72,3.96,'white',radius=True)
text(s,'L3 average tokens',8.13,2.66,4.21,.39,22,'ink',True)
for i,(budget,value) in enumerate([('0 repairs','3,981.6'),('1 repair','6,609.9'),('2 repairs','9,214.1')]):
    y=3.44+i*.71
    text(s,budget,8.16,y,2.18,.38,19,'muted')
    text(s,value,10.33,y,1.99,.4,23,'teal' if i==1 else 'ink',True,PP_ALIGN.RIGHT)
rect(s,8.1,5.72,4.28,.43,'cream',radius=True)
text(s,'More retries ≠ higher success',8.2,5.81,4.08,.3,17,'amber',True,PP_ALIGN.CENTER)
text(s,'Simple L1 tasks: 90% → 80% → 90%; unnecessary repair can misread benign warnings.',.9,6.65,11.8,.28,17,'muted')
note(s,'One repair attempt helps; a second adds cost',65,
     '先说明这是另一个修复预算实验，不能与上一页的 90% 成功率直接拼成同一条提升链。Table II 中，L2 和 L3 都从不修复的 30% 提高到一次修复的 50%；第二次修复仍然是 50%。右侧展示代价：L3 token 从约 3,982 增至 6,610，再到 9,214。因此在这些任务里，一次修复已恢复大部分可恢复案例，再增加预算主要增加开销。L1 则出现 90、80、90 的波动，论文提到有时会把无车辆等良性警告当成失败。',
     'These results also define the limits of what we can claim from this study.',
     'Table II；Sec. IV-D。',
     '论文未充分解释 Table I 与 Table II 各项实验配置的对应关系。主讲只做各表内部比较；答辩前应核对任务子集、初始状态、prompt、修复策略与统计口径。')

# 12 — Limits
s=base('Discussion','What this evidence establishes—and what remains open',
       'A promising workflow result with a bounded empirical scope.', 'Source: Sec. IV–V; reporting limitations identified from the paper')
panel(s,.7,2.5,5.8,3.83,'Supported by this study',
      'Higher L3 execution success\n\nLower tokens and reported TTI\n\nRecovery gains with one repair',body_size=21)
panel(s,6.81,2.5,5.8,3.83,'Open questions',
      '30 tasks; one evaluated backbone\n\nComponents not isolated individually\n\nKPI optimization lacks a quantitative\nconvergence or optimality result',accent='blue',body_size=20)
text(s,'Next: broader benchmarks, component-level ablations, and measured KPI optimization.',.77,6.58,11.9,.36,20,'teal',True)
note(s,'What this evidence establishes—and what remains open',45,
     '这一页主动划清结论边界。证据支持复杂流程的执行可靠性和效率改善，也支持有限修复预算带来恢复收益。但总共只有 30 个任务，实验只评估了一个 LLM backbone。所谓角色消融实际上是单体与完整框架比较，并未逐一隔离 Planner、状态仓库和角色专门化的贡献。双环架构包含 KPI 反馈，但没有报告 KPI 优化的收敛或最优性曲线。页底列的是基于这些局限提出的后续验证建议，不是假装论文已完成的实验。',
     'The main takeaway is a system design principle for reliable traffic simulation workflows.',
     'Sec. IV-A、Table I–II、Sec. V；组件归因和 KPI 定量证据范围是对已报告内容的归纳。')

# 13 — Takeaway
s=base('Takeaway','Reliable automation needs explicit coordination',dark=True,source='Evidence: Table I, complex tasks (L3)')
text(s,'Specialized reasoning  +  persistent artifacts  +  bounded feedback',.69,2.16,12.0,.65,27,'white',True)
vals=[('90%','L3 task success','70% monolithic baseline'),('−47%','L3 token use','11,404.7 → 6,028.7'),('≈5×','L3 TTI speedup','41.4 s → 8.3 s')]
for i,(v,lab,sub) in enumerate(vals):
    x=.7+i*4.17
    rect(s,x,3.36,3.95,2.33,'darkpanel',radius=True)
    text(s,v,x+.25,3.63,3.45,.97,57,'aqua',True)
    text(s,lab,x+.27,4.72,3.4,.39,22,'white',True)
    text(s,sub,x+.27,5.31,3.43,.25,14,'A9BCD0')
text(s,'Natural-language objectives → executable, inspectable SUMO workflows.',.73,6.31,11.95,.61,24,'white')
note(s,'Reliable automation needs explicit coordination',40,
     '用三点收束：专门化角色减少上下文负担，持久化 artifact 保证跨步骤交接，有限反馈帮助恢复执行。最后只重复最关键的 L3 数字：90% 成功率，token 约减少 47%，Time-to-Insight 约五倍加速。再把数字与交通应用连接：目标是让自然语言需求变成可执行、可检查的 SUMO 工作流，降低场景搭建和迭代的门槛。',
     'Thank you. I would be happy to take your questions.',
     'Sec. III；Table I L3；Sec. V。')

# 14 — Closing
s=base('Discussion / Q&A','Thank you',dark=True)
text(s,'Questions?',.74,2.35,10.8,1.16,58,'white',True)
text(s,'Decoupled Intelligence',.78,4.12,10.9,.54,27,'aqua',True)
text(s,'Shuyang Li  ·  Ruimin Ke\nRensselaer Polytechnic Institute',.78,4.99,10.8,.95,22,'white')
text(s,'Supported by the IBM–RPI Future of Computing Research Collaboration.',.79,6.44,11.9,.36,16,'A9BCD0')
note(s,'Thank you / Q&A',10,
     '感谢听众并邀请提问。备用页依次是：B1 完整架构对比结果、B2 完整修复预算结果、B3 角色和文件契约、B4 常见问题与回答边界。若时间已到，停在此页，不继续讲备用页。',
     'Thank you for your attention. I would be happy to take your questions.',
     '论文作者与单位；Acknowledgment。')

# B1 — Full Table I
s=base('Backup B1 / Architecture results','Full results: monolithic vs. multi-agent',
       'Exact values from Table I. ↑ Higher is better; ↓ lower is better.', 'Source: Table I',backup=True)
rows=[['L1 · Basic','Success rate (%) ↑','100.0','90.0'],['','Average tokens ↓','8,529.2','4,845.5'],['','TTI (s) ↓','19.8','6.2'],
      ['L2 · Intermediate','Success rate (%) ↑','100.0','100.0'],['','Average tokens ↓','9,690.3','5,310.1'],['','TTI (s) ↓','22.7','7.3'],
      ['L3 · Complex','Success rate (%) ↑','70.0','90.0'],['','Average tokens ↓','11,404.7','6,028.7'],['','TTI (s) ↓','41.4','8.3']]
table(s,.73,2.5,11.86,['Complexity','Metric','Monolithic','Multi-agent'],rows,[2.8,3.85,2.605,2.605],.385,16)
text(s,'Overall architecture comparison; individual components are not isolated in this table.',.79,6.65,11.77,.3,17,'muted')
note(s,'B1 — Full architecture comparison',0,
     '备用页，按提问定位数据。若有人问简单任务表现，明确回答 L1 成功率是单体 100%、多智能体 90%；不能忽略这一点。若问“每个角色分别贡献多少”，说明该表对比的是完整系统与单体基线，并没有单独去掉某个角色的实验。',
     'This table compares the two complete architectures; it does not isolate each component.',
     'Table I，逐项转录。')

# B2 — Full Table II
s=base('Backup B2 / Repair results','Full results: the repair-budget experiment',
       'Exact values from Table II. Compare repair budgets within this experiment.', 'Source: Table II; Sec. IV-D',backup=True)
rows=[['L1 · Basic','Success rate (%) ↑','90.0','80.0','90.0'],['','Average tokens ↓','4,969.1','5,530.3','6,752.6'],
      ['L2 · Intermediate','Success rate (%) ↑','30.0','50.0','50.0'],['','Average tokens ↓','3,601.6','6,829.5','8,484.8'],
      ['L3 · Complex','Success rate (%) ↑','30.0','50.0','50.0'],['','Average tokens ↓','3,981.6','6,609.9','9,214.1']]
table(s,.73,2.6,11.86,['Complexity','Metric','0 repairs','1 repair','2 repairs'],rows,[2.75,3.05,2.02,2.02,2.02],.47,16)
rect(s,.74,6.18,11.85,.5,'cream',radius=True)
text(s,'No additional L2/L3 success gain from a second attempt; token use continues to rise.',.93,6.28,11.47,.3,17,'amber',True)
note(s,'B2 — Full repair-budget experiment',0,
     '若被问两张表的成功率为何不同，诚实说明论文将其作为两个实验报告，当前稿件没有充分展开配置对应关系；不能未经核对就归因于更难任务、不同 seed 或另一个模型。展示这一表时只比较 0、1、2 次修复预算。另一个重点是最大修复次数，不表示每个任务都实际用了全部次数。',
     'The repair budget is a maximum number of attempts, and these comparisons are within Table II.',
     'Table II；Sec. IV-D。',
     '答辩前优先从原始运行记录补齐两组实验差异，不能现场推测。')

# B3 — Contracts
s=base('Backup B3 / Implementation','What each worker receives and produces',
       'Role contracts and artifact dependencies, simplified from the methodology and prompt appendices.',
       'Source: Sec. III-C; Appendix II–III',backup=True)
rows=[['Builder','City / geographic extent','Build network','.net.xml'],
      ['Modifier','Network + requested edit','Apply topology / TLS edit','Updated .net.xml'],
      ['Demand','Validated network + flow / O–D','Generate matching demand','.rou.xml'],
      ['Runner','Network + routes + run settings','Configure and execute','.sumocfg + outputs'],
      ['Analyst','Simulation outputs + metric','Aggregate KPIs','Metrics + feedback']]
table(s,.72,2.56,11.87,['Worker','Required context','Responsibility','Output'],rows,[1.6,4.0,3.37,2.9],.53,15)
rect(s,.73,6.0,5.72,.71,'white',radius=True)
text(s,'One active role per planned step',.97,6.19,5.25,.35,20,'teal',True)
text(s,'Skip Modifier when no edit is requested',6.94,6.19,5.53,.35,20,'teal',True)
note(s,'B3 — Worker contracts',0,
     '角色输出遵循结构化 action / params 等字段。Planner 每一步仅激活一个 agent_prompt；没有请求修改就跳过 Modifier。这里是简化的输入输出契约，不是完整 API 文档。对多智能体数量的提问，回答一个 Planner 加五种 worker 角色，Orchestrator 管理状态和调度。',
     'The contracts reduce how much each worker has to infer from the conversation history.',
     'Sec. III-C；Appendix II–III。')

# B4 — Q&A boundaries
s=base('Backup B4 / Questions','Answers to likely design questions',source='Source: Sec. III–V; Table I–II',backup=True)
qas=[('Why not a stronger single agent?',
      'This study compares one backbone and one monolithic ReAct design. Broader baselines remain to be tested.'),
     ('Does MCP itself prevent path errors?',
      'MCP standardizes tool access. The Orchestrator manages validated artifact references and dependencies.'),
     ('Is traffic KPI optimization proven?',
      'The feedback mechanism is designed for it. Reported quantitative results focus on execution and repair.'),
     ('Do more agents or retries always help?',
      'No. L1 success is lower in Table I; a second repair adds tokens without improving L2/L3 success in Table II.')]
for i,(q,a) in enumerate(qas):
    y=2.04+i*1.16
    rect(s,.73,y,11.85,1.0,'white',radius=True)
    text(s,q,.94,y+.15,4.22,.73,21,'teal',True)
    text(s,a,5.45,y+.16,6.86,.72,18,'ink')
note(s,'B4 — Likely questions',0,
     '这一页用于准备答辩，不需主动展开。回答采用“本研究比较了什么、证据支持到哪一步”的结构。可以认可更强单体、有状态单体或确定性工作流是值得增加的对照，但不声称已有这些实验。不要把 MCP 说成会自动保证所有正确性；也不要说两次修复无效适用于所有场景。',
     'Our strongest claim is about the reported workflow reliability and efficiency under the evaluated setup.',
     'Sec. III–V；Table I–II。')

output=ROOT/'ITSC_Decoupled_Intelligence_draft.pptx'
prs.save(output)
(ROOT/'source'/'slide_metadata.json').write_text(json.dumps(slides_meta,ensure_ascii=False,indent=2),encoding='utf-8')

md=['# ITSC 现场报告讲稿提示（初稿）','',
    '英文 slides；14 页主报告 + 4 页备用页。主报告目标时长 **12 分 40 秒**，不含提问。',
    '所有结果依据用户提供的 ITSC_ChatSUMO_Multi_agent.pdf。讲稿提示已同步写入 PowerPoint 的“备注”。',
    '主报告按“交通问题 → 失败原因 → 方案 → 证据 → 局限 → 核心结论”组织。','',
    '## 使用方式','',
    '- 排练第一遍：按中文提示用自己的英文表达，重点练习每页的最后一句衔接。',
    '- 需要控制到 12 分钟时：第 3、5、6、12 页各缩短约 10 秒。',
    '- 若时段包含 Q&A：主讲压缩到 10–11 分钟；可以跳过第 3 页和第 6 页，保留结果页。',
    '- 默认结束在第 14 页；第 15–18 页仅用于回答问题。','',
    '## 英文开场参考','',
    '> Good morning. Today I will present Decoupled Intelligence, a multi-agent LLM framework for controllable traffic scenario generation in SUMO. Our question is simple: how can we turn a natural-language traffic request into a reliable, executable simulation workflow? The challenge is keeping the network, demand, configuration, and analysis consistent across many dependent steps.','']
elapsed=0
for m in slides_meta:
    dur=m['seconds']; start=f'{elapsed//60:02}:{elapsed%60:02}'; elapsed+=dur; end=f'{elapsed//60:02}:{elapsed%60:02}'
    md+= [f"## {m['slide']:02d}. {m['title']}",'',f"**时间：{start}–{end}（{dur} 秒）**" if dur else '**备用页，不计入主讲时间。**','',m['zh'],'',f"**英文衔接：** {m['transition']}",'',f"**来源：** {m['source']}",'']
    if m['extra']: md += ['**讲者核对提示：** '+m['extra'],'']
md+=['## 英文收尾参考','',
     '> The main takeaway is that reliable automation needs explicit coordination. Specialized roles reduce the reasoning burden, persistent artifacts keep dependent steps consistent, and bounded feedback helps recover failed executions. In the reported complex-task setting, the framework achieves 90 percent success, uses about 47 percent fewer tokens, and reaches the first valid KPI about five times faster. Thank you for your attention.','']
(ROOT/'speaker_notes_zh.md').write_text('\n'.join(md),encoding='utf-8')

readme='''# ITSC slides 初稿

基于 `/home/chris/Research/ChatSUMO Multi-agent/ITSC_ChatSUMO_Multi_agent.pdf` 制作。
只将论文当作研究材料；附录中的 prompt 约束用于介绍论文的方法，不作为制作 slides 的指令。

## 交付内容

- `ITSC_Decoupled_Intelligence_draft.pptx`：16:9、18 页；14 页主报告 + 4 页备用页。文字、架构图、表格和图表可编辑，图表内嵌数据。
- `ITSC_Decoupled_Intelligence_draft.pdf`：静态预览。
- `speaker_notes_zh.md`：逐页中文讲稿提示、建议时长、英文衔接句和答辩注意点；同样写入 PPT 备注。
- `preview/contact_sheet.jpg`：全套缩略图。
- `source/build_deck.py`：可复现的生成脚本；依赖 python-pptx。
- `source/slide_metadata.json`：逐页结构化备注和来源。

## 设计与范围

- 英文投影页，中文讲者提示，主讲约 12 分 40 秒，不含 Q&A。
- 标题、作者、单位和资助信息沿用论文；没有猜测演讲人、场次、具体日期、会场或代码公开地址。
- 无外部研究或未经确认的新实验；未使用当前代码仓库中较新的实验数据。
- 架构与示例均重画为演示图。街道图是示意，Boston 步骤是说明性示例，不是运行截图或已验证 demo。
- 字体为 Lato；播放设备若无此字体，可安装或统一替换。PDF 可作静态备用。

## 上台前优先核对

1. **Table I 与 Table II 的实验对应关系**：正文未充分描述为何两组 L2/L3 成功率不同。需从原始记录核实任务子集、初始状态、prompt、修复策略和统计口径。初稿只在各表内比较，未将 90% 与 30%→50% 合成一条提升链。
2. **Fig. 3 的单位**：原图同时出现 1.5 km、1.5 miles 和 distance_miles=0.932056788。演示中省略该字段；若加入 demo，先统一范围含义与单位。
3. **TTI 口径**：按正文定义为输入到第一个有效 KPI 的墙钟时间。硬件、服务、缓存、地理数据获取是否计入等细节需要原始运行说明；没有将其描述为纯 SUMO 耗时。
4. **证据范围**：30 个任务，一个被评估的 backbone；无置信区间。Table I 为完整架构对比，不支持单个组件贡献的因果拆分。KPI 双环是方法设计，结果没有量化证明 KPI 收敛或最优性。
5. 确认会议模板、最终演讲人、允许时长是否包含 Q&A；如有真实 demo，可在第 6 页替换为经过验证的短视频或连续截图。

这些核对项不影响审阅初稿；不需要额外信息即可打开和修改。

## 数字核算

- L3 成功率：70% → 90%，增加 **20 个百分点**。
- L3 token：11,404.7 → 6,028.7，减少 **47.14%**。
- L3 TTI：41.4 s → 8.3 s，比值 **4.99×**，时间减少 **79.95%**。
- Table II L2/L3：30% → 50% → 50%；L3 token：3,981.6 → 6,609.9 → 9,214.1。
- Table I L1 成功率：100% → 90%；Table II L1：90% → 80% → 90%。

## 重新生成

```bash
python source/build_deck.py
```

修改脚本后重新生成 PPTX；PDF 需使用 PowerPoint 或 LibreOffice 重新导出。
'''
(ROOT/'README.md').write_text(readme,encoding='utf-8')
print(f'Created {output} ({len(prs.slides)} slides)')
print(f'Main talk: {sum(m["seconds"] for m in slides_meta)} seconds')
